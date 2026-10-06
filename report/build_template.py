"""Build the HECTOR Word report template (docxtpl tags) per TASK spec sections 1-8."""

import os
import re
import sys
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

REPORT_DIR = os.path.dirname(os.path.abspath(__file__))
TEMPLATE_PATH = os.path.join(REPORT_DIR, "templates", "hector_report_template.docx")

HECTOR_VERSION = "v2.1.0"

GOLD = "B8935A"
GOLD_DARK = "9A7634"
INK = "22262E"
MUTE = "7A7F89"
FAINT = "A9ADB5"
LINE = "E4E0D6"

PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
    "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
    "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
    "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind",
    "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection",
    "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
    "sectPr", "pPrChange",
]
RPR_ORDER = [
    "rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike",
    "dstrike", "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
    "vanish", "webHidden", "color", "spacing", "w", "kern", "position", "sz",
    "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText", "vertAlign",
    "rtl", "cs", "em", "lang", "eastAsianLayout", "specVanish", "oMath",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
]
TRPR_ORDER = [
    "cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
    "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
]


def _el(tag, attrs=None):
    el = OxmlElement(tag)
    for k, v in (attrs or {}).items():
        el.set(qn(k), v)
    return el


def insert_ordered(parent, child, order):
    tag = child.tag.split("}")[1]
    idx = order.index(tag)
    for existing in parent:
        etag = existing.tag.split("}")[1]
        if etag in order and order.index(etag) > idx:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


def ensure(parent, tag, order):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        insert_ordered(parent, el, order)
    return el


def set_fonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        insert_ordered(rpr, rf, RPR_ORDER)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + attr), name)
    return rf


def run_props(rpr, font, size_half, color, bold=False, italic=False, spacing=None):
    set_fonts(rpr, font)
    if bold:
        ensure(rpr, "w:b", RPR_ORDER)
        ensure(rpr, "w:bCs", RPR_ORDER)
    if italic:
        ensure(rpr, "w:i", RPR_ORDER)
        ensure(rpr, "w:iCs", RPR_ORDER)
    if color:
        ensure(rpr, "w:color", RPR_ORDER).set(qn("w:val"), color)
    if spacing is not None:
        ensure(rpr, "w:spacing", RPR_ORDER).set(qn("w:val"), str(spacing))
    for t in ("w:sz", "w:szCs"):
        ensure(rpr, t, RPR_ORDER).set(qn("w:val"), str(size_half))


def add_run(par, text, font, size_pt, color=None, bold=False, italic=False,
            spacing=None):
    run = par.add_run(text)
    run_props(run._r.get_or_add_rPr(), font, int(size_pt * 2), color, bold,
              italic, spacing)
    return run


def para_spacing(ppr, before=None, after=None, line=None):
    sp = ensure(ppr, "w:spacing", PPR_ORDER)
    if before is not None:
        sp.set(qn("w:before"), str(before))
    if after is not None:
        sp.set(qn("w:after"), str(after))
    if line is not None:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")
    return sp


def p_border(ppr, edge, sz, space, color):
    pbdr = ensure(ppr, "w:pBdr", PPR_ORDER)
    b = pbdr.find(qn("w:" + edge))
    if b is None:
        b = OxmlElement("w:" + edge)
        pbdr.append(b)
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), color)


def make_style(doc, name, font, size_pt, color, bold=False, italic=False,
               align=None, before=None, after=None, line=None, keep_next=False,
               char_spacing=None, ind_left=None, hanging=None, tab_pos=None,
               numbering=False, border=None):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles["Normal"]
    st.font.name = font
    st.font.size = Pt(size_pt)
    st.font.bold = bold
    st.font.italic = italic
    st.font.color.rgb = RGBColor.from_string(color)
    run_props(st.element.get_or_add_rPr(), font, int(size_pt * 2), color, bold,
              italic, char_spacing)
    ppr = st.element.get_or_add_pPr()
    if align is not None:
        ppr.append(_el("w:jc", {"w:val": align}))
    para_spacing(ppr, before, after, line)
    if keep_next:
        insert_ordered(ppr, _el("w:keepNext"), PPR_ORDER)
    if ind_left is not None:
        ind = _el("w:ind", {"w:left": str(ind_left)})
        if hanging is not None:
            ind.set(qn("w:hanging"), str(hanging))
        insert_ordered(ppr, ind, PPR_ORDER)
    if tab_pos is not None:
        tabs = _el("w:tabs")
        tabs.append(_el("w:tab", {"w:val": "left", "w:pos": str(tab_pos)}))
        insert_ordered(ppr, tabs, PPR_ORDER)
    if border:
        for edge, spec in border.items():
            p_border(ppr, edge, spec[0], spec[1], spec[2])
    if numbering:
        numpr = ppr.get_or_add_numPr()
        numpr.get_or_add_ilvl().val = 0
        numpr.get_or_add_numId().val = 42
    return st


def build_numbering(doc):
    root = doc.part.numbering_part.element
    abs_el = _el("w:abstractNum", {"w:abstractNumId": "42"})
    abs_el.append(_el("w:nsid", {"w:val": "42424242"}))
    abs_el.append(_el("w:multiLevelType", {"w:val": "singleLevel"}))
    lvl = _el("w:lvl", {"w:ilvl": "0"})
    lvl.append(_el("w:start", {"w:val": "1"}))
    lvl.append(_el("w:numFmt", {"w:val": "decimal"}))
    lvl.append(_el("w:suff", {"w:val": "tab"}))
    lvl.append(_el("w:lvlText", {"w:val": "%1"}))
    lvl.append(_el("w:lvlJc", {"w:val": "left"}))
    lvl_ppr = _el("w:pPr")
    tabs = _el("w:tabs")
    tabs.append(_el("w:tab", {"w:val": "num", "w:pos": "440"}))
    lvl_ppr.append(tabs)
    lvl_ppr.append(_el("w:ind", {"w:left": "440", "w:hanging": "440"}))
    lvl.append(lvl_ppr)
    lvl_rpr = _el("w:rPr")
    run_props(lvl_rpr, "Arial", 16, GOLD_DARK, bold=True)
    lvl.append(lvl_rpr)
    abs_el.append(lvl)
    first_num = root.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abs_el)
    else:
        root.append(abs_el)
    num_el = _el("w:num", {"w:numId": "42"})
    num_el.append(_el("w:abstractNumId", {"w:val": "42"}))
    root.append(num_el)


def cell_borders(cell, bottom_sz, bottom_color):
    tcpr = cell._tc.get_or_add_tcPr()
    tcb = ensure(tcpr, "w:tcBorders", TCPR_ORDER)
    for old in list(tcb):
        tcb.remove(old)
    for edge, val in (("top", "nil"), ("left", "nil"),
                      ("bottom", "single"), ("right", "nil")):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), val)
        if edge == "bottom":
            e.set(qn("w:sz"), str(bottom_sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), bottom_color)
        tcb.append(e)


def cell_margins(tbl, top, bottom, left, right):
    mar = ensure(tbl._tbl.tblPr, "w:tblCellMar", TBLPR_ORDER)
    for edge, val in (("top", top), ("bottom", bottom), ("left", left),
                      ("right", right)):
        e = mar.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            mar.append(e)
        e.set(qn("w:w"), str(val))
        e.set(qn("w:type"), "dxa")


def table_borders_none(tbl):
    borders = ensure(tbl._tbl.tblPr, "w:tblBorders", TBLPR_ORDER)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = borders.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            borders.append(e)
        e.set(qn("w:val"), "nil")


def size_table(tbl, widths):
    tblpr = tbl._tbl.tblPr
    tw = ensure(tblpr, "w:tblW", TBLPR_ORDER)
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    tbl.autofit = False
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[i]))
    for row in tbl.rows:
        trpr = row._tr.get_or_add_trPr()
        if trpr.find(qn("w:cantSplit")) is None:
            insert_ordered(trpr, _el("w:cantSplit"), TRPR_ORDER)
        for i, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr()
            tcw = ensure(tcpr, "w:tcW", TCPR_ORDER)
            tcw.set(qn("w:w"), str(widths[i]))
            tcw.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def header_row(tbl, labels):
    row = tbl.rows[0]
    trpr = row._tr.get_or_add_trPr()
    insert_ordered(trpr, _el("w:tblHeader"), TRPR_ORDER)
    for i, cell in enumerate(row.cells):
        cell_borders(cell, 6, GOLD)
        par = cell.paragraphs[0]
        para_spacing(par._p.get_or_add_pPr(), after=0, line=281)
        add_run(par, labels[i], "Arial", 7, GOLD_DARK, bold=True, spacing=40)


def fill_cell(cell, expr, fmt, with_tag=True):
    """marker paragraph + data paragraph ({{ expr }} [+ {{ p.tag }}]) + endfor."""
    marker = cell.paragraphs[0]
    add_run(marker, "{%p for " + expr + " %}", "Georgia", 9, INK)
    data = cell.add_paragraph()
    para_spacing(data._p.get_or_add_pPr(), after=0, line=281)
    add_run(data, "{{ p.text }}", **fmt)
    if with_tag:
        add_run(data, "{{ p.tag }}", "Arial", 7, GOLD_DARK)
    end = cell.add_paragraph()
    add_run(end, "{%p endfor %}", "Georgia", 9, INK)
    return data


def looped_table(doc, widths, labels, row_expr, col_fmts):
    """header row + {%tr for %} marker row + data row + {%tr endfor %} marker row."""
    tbl = doc.add_table(rows=4, cols=3)
    size_table(tbl, widths)
    table_borders_none(tbl)
    cell_margins(tbl, top=100, bottom=100, left=100, right=140)
    header_row(tbl, labels)
    add_run(tbl.rows[1].cells[0].paragraphs[0],
            "{%tr for " + row_expr + " %}", "Georgia", 9, INK)
    data_row = tbl.rows[2]
    for ci, col in enumerate(col_fmts):
        cell_borders(data_row.cells[ci], 4, LINE)
        fill_cell(data_row.cells[ci], "p in r." + col["key"], col["fmt"],
                  with_tag=col.get("tag", True))
    add_run(tbl.rows[3].cells[0].paragraphs[0],
            "{%tr endfor %}", "Georgia", 9, INK)
    return tbl


def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    run_props(normal.element.get_or_add_rPr(), "Georgia", 21, INK)

    make_style(doc, "HectorTitle", "Georgia", 26, INK, bold=True, after=40,
               char_spacing=60)
    make_style(doc, "HectorSectionHeading", "Arial", 8.5, GOLD_DARK, bold=True,
               before=380, after=140, keep_next=True, char_spacing=50)
    make_style(doc, "HectorSubHeading", "Georgia", 9.5, MUTE, italic=True,
               before=200, after=80, keep_next=True)
    make_style(doc, "HectorBody", "Georgia", 10.5, INK, align="both", after=120,
               line=312)
    make_style(doc, "HectorQuery", "Georgia", 12, INK, italic=True, after=120)
    make_style(doc, "HectorNumbered", "Georgia", 10.5, INK, after=100, line=300,
               ind_left=440, hanging=440, numbering=True)
    make_style(doc, "HectorSourceItem", "Georgia", 10, INK, after=80, tab_pos=1100)
    make_style(doc, "HectorTagline", "Georgia", 10, GOLD_DARK, italic=True,
               align="center", before=520, after=140)
    make_style(doc, "HectorDisclaimer", "Georgia", 7, FAINT, italic=True,
               align="center")
    make_style(doc, "HectorFooter", "Arial", 7, FAINT, align="center",
               border={"top": (4, 6, LINE)})

    build_numbering(doc)

    sec = doc.sections[0]
    sectpr = sec._sectPr
    pgsz = sectpr.find(qn("w:pgSz"))
    pgsz.set(qn("w:w"), "11906")
    pgsz.set(qn("w:h"), "16838")
    pgsz.attrib.pop(qn("w:orient"), None)
    pgmar = sectpr.find(qn("w:pgMar"))
    for k, v in (("top", "1200"), ("bottom", "1200"), ("left", "1080"),
                 ("right", "1080"), ("footer", "560"), ("header", "720"),
                 ("gutter", "0")):
        pgmar.set(qn("w:" + k), v)

    p = doc.add_paragraph(style="HectorTitle")
    add_run(p, "H.E.C.T.O.R.", "Georgia", 26, INK, bold=True, spacing=60)

    p = doc.add_paragraph()
    para_spacing(p._p.get_or_add_pPr(), after=0)
    add_run(p, "LEGAL RESEARCH REPORT", "Arial", 7.5, GOLD_DARK, bold=True,
            spacing=70)

    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    para_spacing(ppr, before=60, after=200)
    p_border(ppr, "bottom", 6, 4, GOLD)
    add_run(p, "Hierarchical Evaluation of Civil-Criminal Textual's "
            "Orchestrator & Retrieval", "Georgia", 8.5, FAINT, italic=True)

    meta = doc.add_table(rows=1, cols=3)
    size_table(meta, [3000, 2400, 4346])
    table_borders_none(meta)
    cell_margins(meta, top=0, bottom=0, left=0, right=100)
    for i, (label, tag) in enumerate(
            (("ROUTE", "{{ route }}"), ("CONFIDENCE", "{{ confidence_text }}"),
             ("GENERATED", "{{ generated }}"))):
        cell = meta.cell(0, i)
        lp = cell.paragraphs[0]
        para_spacing(lp._p.get_or_add_pPr(), after=30)
        add_run(lp, label, "Arial", 6.5, FAINT, bold=True, spacing=40)
        vp = cell.add_paragraph()
        para_spacing(vp._p.get_or_add_pPr(), after=0)
        add_run(vp, tag, "Arial", 9, INK)

    p = doc.add_paragraph(style="HectorSectionHeading")
    add_run(p, "QUERY", "Arial", 8.5, GOLD_DARK, bold=True, spacing=50)
    p = doc.add_paragraph(style="HectorQuery")
    add_run(p, "{{ query }}", "Georgia", 12, INK, italic=True)

    p = doc.add_paragraph(style="HectorSectionHeading")
    add_run(p, "RESPONSE", "Arial", 8.5, GOLD_DARK, bold=True, spacing=50)

    p = doc.add_paragraph(style="HectorSubHeading")
    add_run(p, "Direct answer", "Georgia", 9.5, MUTE, italic=True)
    p = doc.add_paragraph(style="HectorBody")
    add_run(p, "{{ direct_answer }}", "Georgia", 10.5, INK)

    p = doc.add_paragraph(style="HectorSubHeading")
    add_run(p, "Statutory text from sources", "Georgia", 9.5, MUTE, italic=True)
    looped_table(
        doc, [1800, 1100, 6846], ["PROVISION", "SOURCE", "TEXT"],
        "r in statutory_rows",
        [
            {"key": "provision_paras",
             "fmt": {"font": "Arial", "size_pt": 8, "color": INK, "bold": True},
             "tag": False},
            {"key": "source_paras",
             "fmt": {"font": "Arial", "size_pt": 7.5, "color": GOLD_DARK},
             "tag": False},
            {"key": "text_paras",
             "fmt": {"font": "Georgia", "size_pt": 9, "color": INK,
                     "italic": True},
             "tag": False},
        ])

    p = doc.add_paragraph(style="HectorSubHeading")
    add_run(p, "Key differences", "Georgia", 9.5, MUTE, italic=True)
    p = doc.add_paragraph()
    add_run(p, "{%p for d in key_differences %}", "Georgia", 10.5, INK)
    p = doc.add_paragraph(style="HectorNumbered")
    add_run(p, "{{ d }}", "Georgia", 10.5, INK)
    p = doc.add_paragraph()
    add_run(p, "{%p endfor %}", "Georgia", 10.5, INK)

    p = doc.add_paragraph(style="HectorSectionHeading")
    add_run(p, "COMPARISON", "Arial", 8.5, GOLD_DARK, bold=True, spacing=50)
    looped_table(
        doc, [2200, 3773, 3773], ["POINT", "IPC", "BNS"],
        "r in comparison_rows",
        [
            {"key": "point_paras",
             "fmt": {"font": "Arial", "size_pt": 8, "color": INK, "bold": True},
             "tag": False},
            {"key": "ipc_paras",
             "fmt": {"font": "Georgia", "size_pt": 9, "color": INK}},
            {"key": "bns_paras",
             "fmt": {"font": "Georgia", "size_pt": 9, "color": INK}},
        ])

    p = doc.add_paragraph(style="HectorSectionHeading")
    add_run(p, "SOURCE SECTIONS", "Arial", 8.5, GOLD_DARK, bold=True, spacing=50)
    p = doc.add_paragraph()
    add_run(p, "{%p for s in sources %}", "Georgia", 10, INK)
    p = doc.add_paragraph(style="HectorSourceItem")
    r = add_run(p, "{{ s.number }}", "Arial", 8, GOLD_DARK, bold=True)
    r._r.append(_el("w:tab"))
    add_run(p, "{{ s.name }}", "Georgia", 10, INK)
    add_run(p, "{{ s.sep }}", "Georgia", 10, INK)
    add_run(p, "{{ s.note_text }}", "Georgia", 9.5, MUTE, italic=True)
    p = doc.add_paragraph()
    add_run(p, "{%p endfor %}", "Georgia", 10, INK)

    p = doc.add_paragraph(style="HectorTagline")
    add_run(p, "Every citation grounded. Every claim verified. Every response "
            "honest.", "Georgia", 10, GOLD_DARK, italic=True)
    p = doc.add_paragraph(style="HectorDisclaimer")
    add_run(p, "This document was generated by HECTOR (Hierarchical Evaluation "
            "of Civil-Criminal Textual's Orchestrator & Retrieval). Responses "
            "are AI-generated and should be verified against official legal "
            "texts. This is not legal advice.", "Georgia", 7, FAINT,
            italic=True)

    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["HectorFooter"]
    para_spacing(fp._p.get_or_add_pPr(), after=0)
    add_run(fp, "HECTOR " + HECTOR_VERSION + "   \u00B7   MIT License   "
            "\u00B7   Hard-RAG Legal Intelligence for India's IPC \u2192 BNS "
            "Transition   \u00B7   ", "Arial", 7, FAINT)
    _field(fp, "PAGE")
    add_run(fp, " / ", "Arial", 7, FAINT)
    _field(fp, "NUMPAGES")

    os.makedirs(os.path.dirname(TEMPLATE_PATH), exist_ok=True)
    doc.save(TEMPLATE_PATH)
    _fix_xml_space(TEMPLATE_PATH)
    print("created:", TEMPLATE_PATH)
    return TEMPLATE_PATH


def _field(par, instr):
    def frun():
        return add_run(par, "", "Arial", 7, FAINT)

    frun()._r.append(_el("w:fldChar", {"w:fldCharType": "begin"}))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " %s " % instr
    frun()._r.append(it)
    frun()._r.append(_el("w:fldChar", {"w:fldCharType": "separate"}))
    add_run(par, "1", "Arial", 7, FAINT)
    frun()._r.append(_el("w:fldChar", {"w:fldCharType": "end"}))


def _fix_xml_space(path):
    tmp = path + ".tmp"
    with zipfile.ZipFile(path, "r") as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                text = data.decode("utf-8")
                text = re.sub(r"<w:t>([^<]*)</w:t>",
                              r'<w:t xml:space="preserve">\1</w:t>', text)
                text = re.sub(
                    r'<w:t (?![^>]*xml:space)([^>]*)>([^<]*)</w:t>',
                    r'<w:t xml:space="preserve" \1>\2</w:t>', text)
                data = text.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


if __name__ == "__main__":
    build()

"""Render the sample payload and assert the rules from spec section 10."""

import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from render_report import render_report  # noqa: E402
from build_template import TEMPLATE_PATH, HECTOR_VERSION  # noqa: E402

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append(bool(ok))
    print("%s - %s%s" % ("PASS" if ok else "FAIL", name,
                         (" | " + str(detail)) if detail else ""))
    return bool(ok)


def collect_texts(doc, footer_xml):
    texts = []
    for p in doc.paragraphs:
        texts.append(p.text)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    texts.append(p.text)
    for p in doc.sections[0].footer.paragraphs:
        texts.append(p.text)
    texts.append(footer_xml)
    return texts


def effective_fonts(doc, footer_xml):
    from docx.oxml import parse_xml
    bad = []
    roots = [doc.element.body, parse_xml(footer_xml.encode("utf-8"))]
    for root in roots:
        for r in root.iter(qn("w:r")):
            ascii_font = None
            rpr = r.find(qn("w:rPr"))
            if rpr is not None:
                rf = rpr.find(qn("w:rFonts"))
                if rf is not None:
                    ascii_font = rf.get(qn("w:ascii"))
            if ascii_font is None:
                ascii_font = "Georgia"
            if ascii_font not in ("Georgia", "Arial"):
                bad.append(ascii_font)
    return bad


def main():
    sample_path = os.path.join(HERE, "sample_data.json")
    with open(sample_path, encoding="utf-8") as fh:
        data = json.load(fh)

    out_dir = os.path.join(tempfile.gettempdir(), "hector_report_test")
    os.makedirs(out_dir, exist_ok=True)
    out_docx = render_report(data, out_dir)

    check("template exists", os.path.exists(TEMPLATE_PATH), TEMPLATE_PATH)
    check("output name matches HECTOR_Report_YYYYMMDD_HHMMSS.docx",
          re.match(r"HECTOR_Report_\d{8}_\d{6}\.docx$",
                   os.path.basename(out_docx)) is not None,
          os.path.basename(out_docx))

    doc = Document(out_docx)
    check("output opens with python-docx", True)

    with zipfile.ZipFile(out_docx) as z:
        doc_xml = z.read("word/document.xml").decode("utf-8")
        footer_xml = z.read("word/footer1.xml").decode("utf-8")

    texts = collect_texts(doc, footer_xml)
    for token in ("|", "**", "---", "{{", "{%"):
        hits = [t[:80] for t in texts if token in t]
        check("no %r anywhere" % token, not hits, hits[:2])

    check("exactly 3 tables (meta + 2)", len(doc.tables) == 3,
          len(doc.tables))
    if len(doc.tables) == 3:
        meta_cells = [c.text.split("\n") for c in doc.tables[0].rows[0].cells]
        check("meta strip labels",
              [m[0] for m in meta_cells] ==
              ["ROUTE", "CONFIDENCE", "GENERATED"],
              [m[0] for m in meta_cells])
        check("meta strip values",
              [m[1] if len(m) > 1 else "" for m in meta_cells] ==
              ["LEGAL_RESEARCH", "92.0%", data["generated"]],
              [m[1] if len(m) > 1 else "" for m in meta_cells])
        hdr1 = [c.text for c in doc.tables[1].rows[0].cells]
        hdr2 = [c.text for c in doc.tables[2].rows[0].cells]
        check("statutory header Provision/Source/Text",
              [h.upper() for h in hdr1] == ["PROVISION", "SOURCE", "TEXT"],
              hdr1)
        check("comparison header Point/IPC/BNS",
              [h.upper() for h in hdr2] == ["POINT", "IPC", "BNS"], hdr2)

    full = "\n".join(texts)
    check("Grounded Answer absent", "Grounded Answer" not in full)
    da = data["direct_answer"].split(".")[0]
    check("direct answer appears exactly once", full.count(da) == 1,
          full.count(da))
    check("confidence printed as 92.0%", "92.0%" in full)
    check("footer contains PAGE field", " PAGE " in footer_xml)
    check("footer contains NUMPAGES field", " NUMPAGES " in footer_xml)
    check("footer carries HECTOR_VERSION", HECTOR_VERSION in footer_xml)

    bad_fonts = effective_fonts(doc, footer_xml)
    check("all runs Georgia or Arial", not bad_fonts, bad_fonts[:5])

    shd = doc_xml.count("<w:shd")
    check("no cell shading", shd == 0, shd)

    if not os.path.exists(SOFFICE):
        check("LibreOffice available", False, SOFFICE)
        print_summary()
        return 1
    check("LibreOffice available", True)

    pdf_name = os.path.splitext(os.path.basename(out_docx))[0] + ".pdf"
    proc = subprocess.run(
        [SOFFICE, "--headless",
         "-env:UserInstallation=file:///C:/Users/DANIEL/AppData/Local/Temp/soffice_profile",
         "--convert-to", "pdf", "--outdir", out_dir, out_docx],
        capture_output=True, text=True, timeout=180)
    pdf_path = os.path.join(out_dir, pdf_name)
    check("PDF conversion succeeded",
          proc.returncode == 0 and os.path.exists(pdf_path),
          (proc.stderr or proc.stdout or "")[-200:])

    if os.path.exists(pdf_path):
        from pypdf import PdfReader
        reader = PdfReader(pdf_path)
        n_pages = len(reader.pages)
        box = reader.pages[0].mediabox
        w_pt, h_pt = float(box.width), float(box.height)
        a4 = abs(w_pt - 595.3) < 5 and abs(h_pt - 841.9) < 5
        check("at most 2 pages", n_pages <= 2, n_pages)
        check("page size A4", a4, "%sx%s pt" % (w_pt, h_pt))

    print_summary(out_docx, pdf_path if os.path.exists(pdf_path) else None)
    return 0 if all(RESULTS) else 1


def print_summary(*paths):
    print("\nFILES CREATED:")
    for p in (TEMPLATE_PATH,) + tuple(p for p in paths if p):
        print("  ", p, "OK" if os.path.exists(p) else "MISSING")
    print("\nASSERTIONS: %d passed, %d failed, %d total" % (
        sum(RESULTS), len(RESULTS) - sum(RESULTS), len(RESULTS)))
    if not all(RESULTS):
        print("RESULT: FAIL")
    else:
        print("RESULT: ALL GREEN")


if __name__ == "__main__":
    sys.exit(main())
